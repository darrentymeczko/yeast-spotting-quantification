"""Selective exports use the committed picks and only publish requested files."""

from dataclasses import replace
from pathlib import Path

import pandas as pd
import pytest
from PIL import Image
from pptx import Presentation

from results_review import export as ex
from results_review import rebuild as rb
from results_review.discovery import SetRun, safe_name
from results_review.model import Candidate, Review
from results_review.rebuild import Frame

REAL_DRAW_FIGURES = ex.draw_figures


def only(**selected):
    values = dict.fromkeys(
        ("data", "summary", "figures", "statistics", "montages", "powerpoint"),
        False)
    values.update(selected)
    return ex.ExportOptions(**values)


@pytest.fixture
def chosen(tmp_path):
    candidates = []
    frames = []
    review = Review(statistical_test="anova", posthoc="tukey", alpha=0.01)
    for medium in ("GLU", "GLY"):
        default = Candidate(medium, medium, "16 Hours", 16, "automatic.JPG",
                            "automatic2.JPG", "least", best_set_score=10)
        pick = replace(default, timepoint="24 Hours", hours=24,
                       plate1=f"{medium}_chosen.JPG", plate2=f"{medium}_chosen2.JPG",
                       extra_plates=(f"{medium}_chosen3.JPG",), best_set_score=1)
        candidates.extend((default, pick))
        review.set_pick(medium, pick)
        review.amend(medium, "rep1", 2, outlier=True)
        tidy = pd.DataFrame([{"experiment": f"Test {medium} 24 Hours",
                              "treatment": f"Test {medium} 24 Hours",
                              "relative_growth": 0.5, "outlier": True,
                              "manual": False, "edit_note": "",
                              "excluded_source": "", "outlier_source": "manual"}])
        frames.append(Frame(medium, pick, f"Test {medium} 24 Hours", tidy, 1))
    run = SetRun(tmp_path / "run", "Test", candidates)
    return run, review, frames


@pytest.fixture
def renderers(monkeypatch):
    seen = {"graphs": [], "montages": []}

    def graph(csv, folder, **statistics):
        seen["graphs"].append((pd.read_csv(csv), statistics))
        out = folder / "figures"
        out.mkdir()
        graphs = []
        for experiment in pd.read_csv(csv).experiment.unique():
            path = out / f"spotting_{safe_name(experiment)}.png"
            Image.new("RGB", (500, 600), "blue").save(path)
            path.with_suffix(".pdf").write_bytes(b"test PDF")
            graphs.append(path)
        stats = out / "spotting_anova.csv"
        stats.write_text("test,p\nanova,0.01\n")
        return graphs, stats

    def montages(root, frames, cfg, folder, cache_dir):
        drawn = {}
        for frame in frames:
            seen["montages"].append(frame.candidate)
            path = folder / "montages" / f"montage_{safe_name(frame.experiment)}.png"
            path.parent.mkdir(exist_ok=True)
            Image.new("RGB", (500, 600), "red" if frame.medium == "GLU" else "green").save(path)
            drawn[frame.medium] = path
        return drawn, {}

    monkeypatch.setattr(ex, "draw_figures", graph)
    monkeypatch.setattr(ex, "draw_montages", montages)
    return seen


def test_powerpoint_only_contains_all_chosen_photos_and_updated_graphs(chosen, renderers):
    run, review, frames = chosen
    result = ex.export(run, review, frames, Path("photos"), {}, options=only(powerpoint=True))
    assert result.ok and not result.warnings
    # Named for the experiment, as the run named the results folder.
    assert [p.name for p in run.chosen_dir.iterdir()] == ["Test.pptx"]
    assert result.csv_path is None and not result.figures and not result.montages
    deck = Presentation(result.powerpoint)
    assert len(deck.slides) == 2
    for slide, frame in zip(deck.slides, frames):
        assert len(slide.shapes) == 2
        assert frame.candidate.plate1 in slide.notes_slide.notes_text_frame.text
        assert frame.candidate.extra_plates[0] in slide.notes_slide.notes_text_frame.text
        assert "automatic.JPG" not in slide.notes_slide.notes_text_frame.text
        assert "tukey" in slide.notes_slide.notes_text_frame.text
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width
            assert shape.top + shape.height <= deck.slide_height
    assert renderers["montages"] == [frame.candidate for frame in frames]
    data, settings = renderers["graphs"][0]
    assert data.outlier.all()
    assert settings == review.statistics_kwargs()


@pytest.mark.parametrize("output,expected", [
    ("data", {ex.CSV_NAME}), ("summary", {ex.SUMMARY_NAME}),
    ("statistics", {"figures/spotting_anova.csv"}),
    ("figures", {"figures/spotting_Test_GLU_24_Hours.png",
                 "figures/spotting_Test_GLU_24_Hours.pdf"}),
    ("montages", {"montages/montage_Test_GLU_24_Hours.png"}),
])
def test_only_requested_output_is_published(chosen, renderers, output, expected):
    run, review, frames = chosen
    result = ex.export(run, review, frames[:1], Path("photos"), {}, options=only(**{output: True}))
    assert result.ok and not result.warnings
    files = {p.relative_to(run.chosen_dir).as_posix()
             for p in run.chosen_dir.rglob("*") if p.is_file()}
    assert files == expected
    if output in ("data", "summary", "montages"):
        assert not renderers["graphs"]
    if output != "montages":
        assert not renderers["montages"]


def test_export_review_rebuilds_only_selected_committed_pick(chosen, renderers, monkeypatch):
    run, review, frames = chosen
    seen = []
    flags = object()

    def rebuild(root, label, cfg, candidate, edits, **kwargs):
        seen.append((candidate, edits, kwargs))
        return frames[1]

    monkeypatch.setattr(rb, "rebuild", rebuild)
    result = ex.export_review(run, review, ["GLY"], Path("photos"), {},
                              options=only(powerpoint=True), outdir=run.chosen_dir,
                              data_flags=flags)
    assert seen == [(frames[1].candidate, review.edits_for("GLY"), {"data_flags": flags})]
    assert len(Presentation(result.powerpoint).slides) == 1
    assert renderers["montages"] == [frames[1].candidate]


def test_failed_graph_never_uses_old_graph_or_overwrites_existing_deck(chosen, renderers, monkeypatch):
    run, review, frames = chosen
    old = run.chosen_dir / "figures" / f"spotting_{safe_name(frames[0].experiment)}.png"
    old.parent.mkdir(parents=True)
    old.write_bytes(b"old graph")
    deck = run.chosen_dir / ex.deck_name(run)
    deck.write_bytes(b"previous deck")
    monkeypatch.setattr(ex, "draw_figures", lambda *a, **k: ([], None))
    result = ex.export(run, review, frames, Path("photos"), {}, options=only(powerpoint=True))
    assert not result.ok and result.powerpoint is None
    assert "no partial PowerPoint" in " ".join(result.warnings)
    assert deck.read_bytes() == b"previous deck"
    assert old.read_bytes() == b"old graph"


def test_rebuild_failure_publishes_no_partial_results(chosen, monkeypatch):
    run, review, frames = chosen

    def rebuild(*args, **kwargs):
        raise rb.RebuildError("missing chosen photo")

    monkeypatch.setattr(rb, "rebuild", rebuild)
    with pytest.raises(rb.RebuildError, match="missing chosen photo"):
        ex.export_review(run, review, run.media, Path("photos"), {},
                         options=only(powerpoint=True), outdir=run.chosen_dir)
    assert not run.chosen_dir.exists()


def test_montages_are_drawn_together_and_a_failure_is_named(chosen, tmp_path,
                                                             monkeypatch):
    """One pass for every medium, as the run draws best/; serial if pools fail."""
    import spotting_timecourse as tc

    run, review, frames = chosen
    monkeypatch.setattr(rb, "pipeline_candidate", lambda root, cand, cfg: cand)
    monkeypatch.setattr(rb, "rows_for", lambda cand, cfg=None: (0,))
    monkeypatch.setattr(rb, "layout_for", lambda cfg: None)
    monkeypatch.setattr(tc, "montage_job",
                        lambda cand, rows, strains, out, cache, experiment, layout:
                        (experiment, Path(out)))
    passes = []

    def draw(jobs, workers):
        passes.append((len(jobs), workers))
        if workers > 1:
            raise OSError("process pools are blocked on this computer")
        experiment, out = jobs[0]
        out.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGB", (5, 5)).save(out)
        return [(jobs[1][0], "RuntimeError: photo unreadable")]

    monkeypatch.setattr(tc, "draw_montages", draw)
    drawn, failed = ex.draw_montages(Path("photos"), frames, {"strains": []},
                                     tmp_path / "work", workers=8)
    # Never more workers than montages; one retry, serially, when pools fail.
    assert passes == [(2, 2), (2, 1)]
    assert set(drawn) == {"GLU"} and drawn["GLU"].exists()
    assert failed == {"GLY": "RuntimeError: photo unreadable"}


def _pipeline_sheets(run, frames):
    """Each chosen candidate's comparison sheet: black spots beside a red graph."""
    import spotting_sheet as ss

    for frame in frames:
        pieces = {"montage": Image.new("RGB", (900, 400), (0, 0, 0)),
                  "graph": Image.new("RGB", (800, 400), (200, 30, 30))}
        im, lay = ss.arrange(pieces, ss.SHEET_BOX)
        ss.save(im, lay, run.sheet(frame.candidate), panels=pieces)


def _slide_pictures(deck_path):
    import io

    deck = Presentation(deck_path)
    return [[Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
             for shape in slide.shapes] for slide in deck.slides]


@pytest.mark.parametrize("view,rotate", [("fit", False), ("aligned", False),
                                         ("fit", True), ("aligned", True)])
def test_slides_show_each_sheet_the_way_the_review_does(chosen, renderers,
                                                        monkeypatch, view, rotate):
    import spotting_sheet as ss

    run, review, frames = chosen
    _pipeline_sheets(run, frames)
    asked, real_view = [], ss.view

    def recorded(sheet, mode, box, rotate=False):
        asked.append((mode, rotate))
        return real_view(sheet, mode, box, rotate)

    monkeypatch.setattr(ss, "view", recorded)
    result = ex.export(run, review, frames, Path("photos"), {},
                       options=only(powerpoint=True, view=view, rotate=rotate))
    assert result.powerpoint is not None
    # The review's own view of each sheet was asked for, slide by slide.
    assert asked[0] == (view, rotate)
    # One picture per slide -- the sheet laid out -- and no montage redrawn
    # from the photographs to make it.
    pictures = _slide_pictures(result.powerpoint)
    assert [len(p) for p in pictures] == [1, 1] and renderers["montages"] == []
    for (picture,) in pictures:
        colours = {c for _, c in picture.getcolors(1 << 20)}
        assert (0, 0, 0) in colours                    # the sheet's spots
        assert (0, 0, 255) in colours                  # the corrected graph
        assert (200, 30, 30) not in colours            # never the pipeline's
    # These sheets carry no sideways graph or spot positions, so anything but
    # plain Fit falls back to it -- and says so rather than failing the deck.
    fell_back = [w for w in result.warnings if "shown as Fit" in w]
    assert len(fell_back) == (0 if (view, rotate) == ("fit", False) else 2)


def test_a_candidate_without_a_sheet_keeps_montage_and_graph(chosen, renderers):
    run, review, frames = chosen
    _pipeline_sheets(run, frames[:1])
    result = ex.export(run, review, frames, Path("photos"), {},
                       options=only(powerpoint=True, view="aligned"))
    assert result.ok
    assert [len(p) for p in _slide_pictures(result.powerpoint)] == [1, 2]
    assert renderers["montages"] == [frames[1].candidate]


def test_exports_name_treatments_as_the_experiment_does(chosen, renderers):
    """Graph titles, files and rows say "K-OAc NaAsO2", never "K-OACNAA"."""
    run, review, frames = chosen
    labelled = [replace(c, medium_label=f"{c.medium} as named")
                for c in run.candidates]
    run.candidates[:] = labelled
    result = ex.export(run, review, frames, Path("photos"), {},
                       options=only(data=True, figures=True, montages=True,
                                    powerpoint=True))
    assert result.ok
    data = pd.read_csv(result.csv_path, encoding="utf-8-sig")
    assert set(data.experiment) == set(data.treatment) == {
        "Test GLU as named 24 Hours", "Test GLY as named 24 Hours"}
    # The renderer titles each graph and names its file from that string.
    graphs, _ = renderers["graphs"][0]
    assert set(graphs.treatment) == set(data.treatment)
    names = {p.name for p in run.chosen_dir.rglob("*.png")}
    assert {"spotting_Test_GLU_as_named_24_Hours.png",
            "montage_Test_GLU_as_named_24_Hours.png"} <= names
    deck = Presentation(result.powerpoint)
    assert "Test GLY as named 24 Hours" in \
        deck.slides[1].notes_slide.notes_text_frame.text
    # The review's own frames are not renamed by exporting them.
    assert frames[0].experiment == "Test GLU 24 Hours"


def test_names_that_would_share_a_file_keep_their_codes(chosen):
    run, _, frames = chosen
    run.candidates[:] = [replace(c, medium_label="Glucose + 37C"
                                 if c.medium == "GLU" else "Glucose 37C")
                         for c in run.candidates]
    assert [f.experiment for f in ex.named_frames(run, frames)] == [
        "Test GLU 24 Hours", "Test GLY 24 Hours"]


@pytest.fixture
def held():
    """Hold a file open the way PowerPoint holds a deck: no sharing at all."""
    import sys

    if sys.platform != "win32":
        pytest.skip("Windows file locking")
    import ctypes
    from ctypes import wintypes

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateFileW.restype = wintypes.HANDLE
    kernel32.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD,
                                     wintypes.DWORD, wintypes.LPVOID,
                                     wintypes.DWORD, wintypes.DWORD,
                                     wintypes.HANDLE]
    handles = []

    def hold(path):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(b"the deck somebody has open")
        handle = kernel32.CreateFileW(str(path), 0x80000000 | 0x40000000, 0,
                                      None, 3, 0x80, None)
        assert handle not in (None, wintypes.HANDLE(-1).value)
        handles.append(handle)
        return path

    yield hold
    for handle in handles:
        kernel32.CloseHandle(handle)


def test_an_open_deck_is_reported_before_any_work(chosen, renderers, held):
    run, review, frames = chosen
    deck = held(run.chosen_dir / ex.deck_name(run))
    assert ex.in_use([deck, run.chosen_dir / "absent.csv"]) == [deck]
    with pytest.raises(ValueError, match=r"Test\.pptx is open in another program"):
        ex.export(run, review, frames, Path("photos"), {},
                  options=only(powerpoint=True))
    assert not renderers["graphs"] and not renderers["montages"]


def test_a_deck_opened_mid_export_is_named_not_a_winerror(chosen, renderers,
                                                          held, monkeypatch):
    run, review, frames = chosen
    held(run.chosen_dir / ex.deck_name(run))
    monkeypatch.setattr(ex, "in_use", lambda paths: [])   # opened after the check
    result = ex.export(run, review, frames, Path("photos"), {},
                       options=only(powerpoint=True))
    assert result.powerpoint is None
    assert any("Test.pptx is open in another program" in w
               for w in result.warnings)
    assert not any("WinError" in w for w in result.warnings)


def test_dialog_says_the_deck_is_open_and_stays_open(chosen, tk_root, held,
                                                     monkeypatch):
    from results_review.gui import export_dialog as ui
    run, review, _ = chosen
    held(run.chosen_dir / ex.deck_name(run))
    errors = []
    monkeypatch.setattr(ui.messagebox, "showerror", lambda *a, **k: errors.append(a))
    dialog = ui.ExportDialog(tk_root, run, review)
    dialog._deck_only()
    dialog._accept()
    assert errors and "Test.pptx is open" in errors[0][1]
    assert dialog.result is None and dialog.winfo_exists()
    dialog.destroy()


@pytest.mark.parametrize("label,expected", [
    ("Set13", "Set13.pptx"),
    ("Set 13 - Take 2", "Set 13 - Take 2.pptx"),
    ("Glucose: 37C/NaAsO2?", "Glucose_ 37C_NaAsO2_.pptx"),
])
def test_the_deck_is_named_for_its_experiment(tmp_path, label, expected):
    assert ex.deck_name(SetRun(tmp_path, label)) == expected


def test_unselected_existing_files_are_preserved(chosen, renderers):
    run, review, frames = chosen
    run.chosen_dir.mkdir(parents=True)
    old = run.chosen_dir / ex.CSV_NAME
    old.write_bytes(b"keep data")
    result = ex.export(run, review, frames, Path("photos"), {}, options=only(powerpoint=True))
    assert result.ok and old.read_bytes() == b"keep data"


def test_no_outputs_and_pipeline_destinations_are_rejected(chosen):
    run, review, frames = chosen
    with pytest.raises(ValueError, match="at least one output"):
        ex.export(run, review, frames, Path("photos"), {}, options=only())
    for folder in (run.best_dir, run.best_dir / "nested", run.results_dir):
        with pytest.raises(ValueError, match="outside best"):
            ex.export(run, review, frames, Path("photos"), {}, outdir=folder)


def test_dialog_selects_powerpoint_only_and_a_treatment(chosen, tk_root):
    from results_review.gui.export_dialog import ExportDialog
    run, review, _ = chosen
    dialog = ExportDialog(tk_root, run, review)
    assert len(dialog.treatments.curselection()) == 2
    dialog._deck_only()
    dialog.treatments.selection_clear(0, "end")
    dialog.treatments.selection_set(1)
    dialog._accept()
    assert dialog.result.options == only(powerpoint=True)
    assert dialog.result.media == ["GLY"]
    assert dialog.result.outdir == run.chosen_dir.resolve()


def test_dialog_carries_the_reviews_view_to_the_slides(chosen, tk_root):
    from results_review.gui.export_dialog import ExportDialog
    run, review, _ = chosen
    dialog = ExportDialog(tk_root, run, review, "aligned", True)
    dialog._deck_only()
    dialog._accept()
    assert dialog.result.options == only(powerpoint=True, view="aligned",
                                          rotate=True)


def test_dialog_rejects_empty_selection_and_cancels_cleanly(chosen, tk_root, monkeypatch):
    from results_review.gui import export_dialog as ui
    run, review, _ = chosen
    errors = []
    monkeypatch.setattr(ui.messagebox, "showerror", lambda *a, **k: errors.append(a))
    dialog = ui.ExportDialog(tk_root, run, review)
    dialog.treatments.selection_clear(0, "end")
    dialog._accept()
    assert errors and dialog.result is None
    assert dialog.winfo_exists()
    dialog.destroy()
    assert not run.chosen_dir.exists()


def test_powerpoint_with_real_updated_graphs(chosen, renderers, monkeypatch):
    run, review, frames = chosen
    review.statistical_test = "t_test"
    review.p_adjust = "holm"
    monkeypatch.setattr(ex, "draw_figures", REAL_DRAW_FIGURES)
    for frame in frames:
        rows = []
        for rep in range(4):
            for col, strain in ((1, "Control"), (2, "Mutant")):
                raw = 10 + rep if col == 1 else 4 + rep
                rows.append(dict(experiment=frame.experiment, treatment=frame.experiment,
                                 set="TC", plate=1 + rep // 2, image="chosen.JPG",
                                 replicate=f"rep{rep}", dilution_row=1, dilution="least",
                                 strain_col=col, strain=strain, raw_growth=raw,
                                 artifact=False, excluded=False, is_control=col == 1,
                                 control_raw=10 + rep, control_mean=10 + rep,
                                 relative_growth=raw / (10 + rep), control_ok=True,
                                 outlier=rep == 0 and col == 2))
        frame.tidy = pd.DataFrame(rows)
    result = ex.export(run, review, frames, Path("photos"), {}, options=only(powerpoint=True))
    assert result.ok and not result.warnings
    deck = Presentation(result.powerpoint)
    assert len(deck.slides) == 2
    for slide in deck.slides:
        assert slide.shapes[1].image.size[0] > 500
        assert "holm" in slide.notes_slide.notes_text_frame.text


@pytest.mark.parametrize("cancel", [False, True])
def test_export_button_schedules_a_fixed_snapshot_or_cancels(chosen, monkeypatch, cancel):
    from types import SimpleNamespace
    from results_review.app import ReviewApp
    from results_review.gui import export_dialog as ui

    run, review, frames = chosen
    request = ui.ExportRequest(only(powerpoint=True), ["GLY"], run.chosen_dir)
    asked = []
    monkeypatch.setattr(ui, "ask_export", lambda *args: (
        asked.append(args), None if cancel else request)[1])
    calls, jobs = [], []
    ctl = SimpleNamespace(can_edit=True, run=run, review=review,
                          capture_root=Path("photos"), cfg={"strains": ["Control"]},
                          save=lambda: calls.append("save"), data_flags=lambda: None)
    app = SimpleNamespace(_busy=False, ctl=ctl, winfo_toplevel=lambda: None,
                          sheet=SimpleNamespace(view="aligned", rotated=True),
                          _run_async=lambda work, done, text: jobs.append(work),
                          _exported=lambda result: None, _progress=lambda message: None)
    exported = []
    monkeypatch.setattr(ex, "export_review", lambda *a, **k: exported.append((a, k)))
    ReviewApp._export(app)
    # The dialog is told how the sheet is being shown, for the slides.
    assert asked[0][3:] == ("aligned", True)
    if cancel:
        assert not calls and not jobs
        return
    assert calls == ["save"] and len(jobs) == 1
    review.set_pick("GLY", run.pipeline_best("GLY"))
    review.alpha = 0.2
    ctl.cfg["strains"].append("Later edit")
    jobs[0]()
    args, kwargs = exported[0]
    assert args[1].pick("GLY").matches(frames[1].candidate)
    assert args[1].alpha == 0.01
    assert args[2] == ["GLY"]
    assert args[4] == {"strains": ["Control"]}
    assert kwargs["options"] == only(powerpoint=True)
